from base64 import b64decode
from pathlib import Path
import zipfile

import pytest
from lxml import etree
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# A valid one-pixel PNG; it is test data written only under pytest's tmp_path.
PNG_1X1 = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4////fwAJ+wP+NfZq7AAAAABJRU5ErkJggg==")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


@pytest.fixture()
def minimal_docx(tmp_path: Path) -> Path:
    destination = tmp_path / "minimal.docx"
    document = Document()
    document.add_paragraph("A minimal document.")
    document.save(destination)
    return destination


@pytest.fixture()
def rich_docx(tmp_path: Path) -> Path:
    image = tmp_path / "pixel.png"
    image.write_bytes(PNG_1X1)
    destination = tmp_path / "rich.docx"
    document = Document()
    document.add_heading("A real heading", level=1)
    paragraph = document.add_paragraph("Plain paragraph with ")
    add_hyperlink(paragraph, "a hyperlink", "https://example.com")
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = "Right"
    document.add_picture(str(image))
    document.save(destination)
    return destination


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _replace_zip_parts(docx_path: Path, replacements: dict[str, bytes]) -> None:
    rebuilt = docx_path.with_suffix(".rebuilt.docx")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            if info.filename not in replacements:
                destination.writestr(info, source.read(info.filename))
        for name, data in replacements.items():
            destination.writestr(name, data)
    rebuilt.replace(docx_path)


def _xml(value: str) -> etree._Element:
    return etree.fromstring(value.encode("utf-8"))


@pytest.fixture()
def m2_docx(tmp_path: Path) -> Path:
    """Test-only OOXML fixture containing constructs python-docx cannot author."""
    destination = tmp_path / "m2-constructs.docx"
    document = Document()
    document.add_heading("M2 heading", level=1)
    paragraph = document.add_paragraph("Linked ")
    add_hyperlink(paragraph, "visible link", "https://example.com/linked")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell text"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header text"
    section.footer.paragraphs[0].text = "Footer text"
    document.save(destination)

    with zipfile.ZipFile(destination, "r") as package:
        root = etree.fromstring(package.read("word/document.xml"))
    body = root.find(f"{{{W_NS}}}body")
    insertion_index = len(body) - 1  # preserve w:sectPr as final body child
    raw_constructs = [
        """<w:sdt xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:sdtPr><w:tag w:val='fixture-control'/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>Controlled text</w:t></w:r></w:p></w:sdtContent></w:sdt>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> UNKNOWNFIELD target </w:instrText></w:r><w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>Field result</w:t></w:r><w:r><w:fldChar w:fldCharType='end'/></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:fldSimple w:instr=' PAGE '><w:r><w:t>7</w:t></w:r></w:fldSimple></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:del><w:r><w:delText>Deleted text</w:delText></w:r></w:del></w:p>""",
        """<mc:AlternateContent xmlns:mc='http://schemas.openxmlformats.org/markup-compatibility/2006' xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><mc:Choice Requires='w14'><w:p><w:r><w:t>Choice text</w:t></w:r></w:p></mc:Choice><mc:Fallback><w:p><w:r><w:t>Fallback text</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>""",
        """<mc:AlternateContent xmlns:mc='http://schemas.openxmlformats.org/markup-compatibility/2006' xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><mc:Choice Requires='w14'><w:p><w:r><w:t>Same alternate text</w:t></w:r></w:p></mc:Choice><mc:Fallback><w:p><w:r><w:t>Same alternate text</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>""",
        """<mc:AlternateContent xmlns:mc='http://schemas.openxmlformats.org/markup-compatibility/2006' xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><mc:Choice Requires='w14'><w:p><w:r><w:t>Multi </w:t></w:r><w:r><w:t>alternate text</w:t></w:r></w:p></mc:Choice><mc:Fallback><w:p><w:r><w:t>Multi alternate text</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing' xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'><w:r><w:drawing><wp:inline><a:graphic><a:graphicData><w:txbxContent><w:p><w:r><w:t>Drawing textbox text</w:t></w:r></w:p></w:txbxContent></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' xmlns:v='urn:schemas-microsoft-com:vml'><w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>VML textbox text</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>""",
    ]
    for raw in raw_constructs:
        body.insert(insertion_index, _xml(raw))
        insertion_index += 1
    footnotes = """<w:footnotes xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:footnote w:id='17'><w:p><w:r><w:t>Footnote text</w:t></w:r></w:p></w:footnote></w:footnotes>"""
    endnotes = """<w:endnotes xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:endnote w:id='18'><w:p><w:r><w:t>Endnote text</w:t></w:r></w:p></w:endnote></w:endnotes>"""
    comments = """<w:comments xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:comment w:id='1'><w:p><w:r><w:t>Comment text</w:t></w:r></w:p></w:comment></w:comments>"""
    _replace_zip_parts(destination, {
        "word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
        "word/footnotes.xml": footnotes.encode(),
        "word/endnotes.xml": endnotes.encode(),
        "word/comments.xml": comments.encode(),
    })
    return destination


@pytest.fixture()
def unknown_shape_docx(minimal_docx: Path) -> Path:
    with zipfile.ZipFile(minimal_docx, "r") as package:
        root = etree.fromstring(package.read("word/document.xml"))
    body = root.find(f"{{{W_NS}}}body")
    body.insert(len(body) - 1, _xml("""<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'><w:r><w:drawing><wp:inline/></w:drawing></w:r></w:p>"""))
    _replace_zip_parts(minimal_docx, {"word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)})
    return minimal_docx


@pytest.fixture()
def identity_docx(tmp_path: Path) -> Path:
    """A coverage-complete M3 fixture with every supported writable story type."""
    image = tmp_path / "identity-pixel.png"
    image.write_bytes(PNG_1X1)
    destination = tmp_path / "identity.docx"
    document = Document()
    document.add_heading("Identity heading", level=1)
    paragraph = document.add_paragraph("Identity ")
    add_hyperlink(paragraph, "hyperlink", "https://example.com/identity")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Identity cell"
    document.add_picture(str(image))
    section = document.sections[0]
    section.header.paragraphs[0].text = "Identity header"
    section.footer.paragraphs[0].text = "Identity footer"
    document.save(destination)
    with zipfile.ZipFile(destination, "r") as package:
        root = etree.fromstring(package.read("word/document.xml"))
    body = root.find(f"{{{W_NS}}}body")
    insertion_index = len(body) - 1
    raw_constructs = [
        """<w:sdt xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:sdtContent><w:p><w:r><w:t>Identity controlled</w:t></w:r></w:p></w:sdtContent></w:sdt>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:fldSimple w:instr=' PAGE '><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing' xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'><w:r><w:drawing><wp:inline><a:graphic><a:graphicData><w:txbxContent><w:p><w:r><w:t>Identity drawing textbox</w:t></w:r></w:p></w:txbxContent></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' xmlns:v='urn:schemas-microsoft-com:vml'><w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Identity VML textbox</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>""",
    ]
    for raw in raw_constructs:
        body.insert(insertion_index, _xml(raw))
        insertion_index += 1
    _replace_zip_parts(destination, {
        "word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })
    return destination


@pytest.fixture()
def tagged_docx(tmp_path: Path) -> Path:
    """M4.7/M4.8 fixture with semantic anchors and structural-island cases."""
    destination = tmp_path / "tagged.docx"
    image = tmp_path / "tagged-pixel.png"
    image.write_bytes(PNG_1X1)
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Only ")
    bold = paragraph.add_run("under these circumstances")
    bold.bold = True
    paragraph.add_run(" can it work.")
    italic_paragraph = document.add_paragraph("This is ")
    italic = italic_paragraph.add_run("important")
    italic.italic = True
    hyperlink_paragraph = document.add_paragraph("Read ")
    add_hyperlink(hyperlink_paragraph, "this article", "https://example.com/article")
    document.add_paragraph("Visit ") .add_run("https://example.com/protected")
    binary_drawing = document.add_paragraph()
    binary_drawing.add_run("Binary ")
    binary_drawing.add_run().add_picture(str(image))
    binary_drawing.add_run(" image")
    document.save(destination)
    with zipfile.ZipFile(destination, "r") as package:
        root = etree.fromstring(package.read("word/document.xml"))
    protected_text = root.xpath(".//w:t[text()='https://example.com/protected']", namespaces={"w": W_NS})[0]
    protected_run = protected_text.getparent()
    protected_properties = etree.Element(f"{{{W_NS}}}rPr")
    etree.SubElement(protected_properties, f"{{{W_NS}}}spacing", {f"{{{W_NS}}}val": "40"})
    etree.SubElement(protected_properties, f"{{{W_NS}}}w", {f"{{{W_NS}}}val": "80"})
    protected_run.insert(0, protected_properties)
    body = root.find(f"{{{W_NS}}}body")
    raw_constructs = [
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:fldSimple w:instr=' PAGE '><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing' xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'><w:r><w:t>Caption </w:t></w:r><w:r><w:drawing><wp:inline><a:graphic><a:graphicData><w:txbxContent><w:p><w:r><w:t>Nested textbox</w:t></w:r></w:p></w:txbxContent></a:graphicData></a:graphic></wp:inline></w:drawing></w:r><w:r><w:t> after image</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Before </w:t></w:r><w:bookmarkStart w:id='77' w:name='keep-bookmark'/><w:r><w:t>After</w:t></w:r><w:bookmarkEnd w:id='77'/></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Left</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>Right</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Unsafe mixed</w:t><w:tab/></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:rPr><w:spacing w:val='40'/><w:w w:val='80'/></w:rPr><w:t>Micro layout English</w:t></w:r></w:p>""",
    ]
    for raw in raw_constructs:
        body.insert(len(body) - 1, _xml(raw))
    _replace_zip_parts(destination, {"word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)})
    return destination


@pytest.fixture()
def format_audit_docx(tmp_path: Path) -> Path:
    """Raw OOXML fixture for M4.5 direct-format and fragment evidence."""
    destination = tmp_path / "format-audit.docx"
    document = Document()
    document.add_heading("Short Title", level=1)
    document.add_paragraph("Previous complete sentence.")
    document.add_paragraph("friends,")
    document.add_paragraph("Next complete sentence.")
    document.save(destination)
    with zipfile.ZipFile(destination, "r") as package:
        root = etree.fromstring(package.read("word/document.xml"))
    body = root.find(f"{{{W_NS}}}body")
    insertion_index = len(body) - 1
    raw_constructs = [
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:rPr><w:lang w:val='en-US'/></w:rPr><w:t>Hello</w:t></w:r><w:r><w:rPr><w:lang w:val='en-GB'/></w:rPr><w:t> world</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Plain </w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>bold</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Plain </w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>italic</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:rPr/><w:t>Normal </w:t></w:r><w:r><w:rPr><w:vertAlign w:val='superscript'/></w:rPr><w:t>super</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Intro </w:t></w:r><w:r><w:t>https://example.com</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:pPr><w:numPr><w:ilvl w:val='0'/><w:numId w:val='1'/></w:numPr></w:pPr><w:r><w:t>item</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>This sentence continues</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>across blocks.</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>perhaps</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Final sentence.</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>Before fragment</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>broken,</w:t></w:r></w:p>""",
        """<w:p xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:r><w:t>continues after.</w:t></w:r></w:p>""",
    ]
    for raw in raw_constructs:
        body.insert(insertion_index, _xml(raw))
        insertion_index += 1
    _replace_zip_parts(destination, {"word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)})
    return destination
