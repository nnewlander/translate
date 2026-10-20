"""Build the deterministic, book-neutral DOCX fixture used by the core smoke test."""

from __future__ import annotations

from base64 import b64decode
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import zipfile


PNG_1X1 = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4////fwAJ+wP+NfZq7AAAAABJRU5ErkJggg=="
)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build(destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    image = destination.with_suffix(".png")
    image.write_bytes(PNG_1X1)
    try:
        document = Document()
        document.add_heading("Generic Pipeline Smoke Book", level=0)
        document.add_heading("A Heading", level=1)
        paragraph = document.add_paragraph("A plain paragraph with ")
        _add_hyperlink(paragraph, "a hyperlink", "https://example.com/smoke")
        document.add_paragraph("First list item", style="List Bullet")
        document.add_paragraph("Second list item", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Column A"
        table.cell(0, 1).text = "Column B"
        table.cell(1, 0).text = "Cell one"
        table.cell(1, 1).text = "Cell two"
        document.add_picture(str(image))
        note_paragraph = document.add_paragraph("A paragraph with a native footnote")
        note_run = note_paragraph.add_run()
        note_reference = OxmlElement("w:footnoteReference")
        note_reference.set(qn("w:id"), "1")
        note_run._r.append(note_reference)
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("Second section paragraph")
        document.save(destination)
        _add_native_footnote_part(destination)
    finally:
        image.unlink(missing_ok=True)


def _add_native_footnote_part(destination: Path) -> None:
    """Add the minimal Word footnotes part, relationship, and content type."""
    with zipfile.ZipFile(destination, "r") as package:
        replacements = {
            "word/footnotes.xml": b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>Native footnote text</w:t></w:r></w:p></w:footnote>
</w:footnotes>''',
        }
        rels = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        etree.SubElement(rels, f"{{{rel_ns}}}Relationship", {
            "Id": "rIdGenericSmokeFootnotes",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
            "Target": "footnotes.xml",
        })
        replacements["word/_rels/document.xml.rels"] = etree.tostring(
            rels, xml_declaration=True, encoding="UTF-8", standalone=True,
        )
        content_types = etree.fromstring(package.read("[Content_Types].xml"))
        ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        etree.SubElement(content_types, f"{{{ct_ns}}}Override", {
            "PartName": "/word/footnotes.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        })
        replacements["[Content_Types].xml"] = etree.tostring(
            content_types, xml_declaration=True, encoding="UTF-8", standalone=True,
        )
        rebuilt = destination.with_suffix(".rebuilt.docx")
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as output:
            for info in package.infolist():
                if info.filename not in replacements:
                    output.writestr(info, package.read(info.filename))
            for name, payload in replacements.items():
                output.writestr(name, payload)
    rebuilt.replace(destination)


if __name__ == "__main__":
    build(Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "generic_pipeline_smoke.docx")
